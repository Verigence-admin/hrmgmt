from __future__ import annotations

import io
from collections import defaultdict
from datetime import date, timedelta
from typing import Annotated, Any

from fastapi import APIRouter, Depends, Query, Request
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import Connection, text

from hrmgmt import permissions as perm
from hrmgmt import workcontext as wc
from hrmgmt.api.attendance import day_kind, get_clock
from hrmgmt.attendance_report import DELINQUENCY_LABELS, MAX_RANGE_DAYS, Row, build_rows, row_view
from hrmgmt.audit import record_audit
from hrmgmt.db import get_conn
from hrmgmt.errors import ApiError
from hrmgmt.principal import require_permission
from hrmgmt.security import HumanPrincipal
from hrmgmt.timeutil import Clock, ist_date, to_ist

router = APIRouter(prefix="/hr/v1", tags=["Attendance reports"])

can_read_attendance = require_permission(perm.HR_ATTENDANCE_READ_ALL)
can_read_employees = require_permission(perm.HR_EMPLOYEE_READ)


def _day(value: str | None, default: date, name: str) -> date:
    if not value:
        return default
    try:
        return date.fromisoformat(value)
    except ValueError as exc:
        raise ApiError(422, "HR_VALIDATION_FAILED", f"{name} must look like 2026-10-05.") from exc


@router.get("/attendance/daily")
def daily_attendance(
    on: Annotated[str | None, Query(alias="date", max_length=10)] = None,
    project_code: Annotated[str | None, Query(alias="projectCode", max_length=60)] = None,
    _: HumanPrincipal = Depends(can_read_attendance),
    clock: Clock = Depends(get_clock),
    conn: Connection = Depends(get_conn),
) -> dict[str, Any]:
    """HR and the CEO: every active employee's check-in and check-out on one day, project by
    project, with whatever needs attention."""
    today = ist_date(clock())
    day = _day(on, today, "Date")
    if day > today:
        raise ApiError(422, "HR_VALIDATION_FAILED", "A future day has no attendance yet.")
    rows = build_rows(conn, day, day, today=today, project_code=project_code)
    kind, holiday_name = day_kind(conn, day)
    people: dict[str, Row] = {}
    for r in rows:
        previous = people.get(r.employee_id)
        if previous is None or (not previous.delinquencies and r.delinquencies):
            people[r.employee_id] = r
    counts = [r.status for r in people.values()]
    return {
        "date": day.isoformat(),
        "dayKind": kind,
        "holiday": holiday_name,
        "summary": {
            "employees": len(people),
            "checkedIn": sum(1 for r in people.values() if r.check_in_at),
            "completed": counts.count("COMPLETE"),
            "stillIn": counts.count("CHECKED_IN"),
            "notCheckedIn": counts.count("NOT_CHECKED_IN"),
            "absent": counts.count("ABSENT"),
            "onLeave": counts.count("ON_LEAVE"),
            "pendingApproval": counts.count("PENDING_APPROVAL"),
            "withDelinquencies": sum(1 for r in people.values() if r.delinquencies),
        },
        "rows": [row_view(r) for r in rows],
    }


def _text(value: Any) -> Any:
    """A spreadsheet cell that starts with = + - or @ would be read as a formula."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        return "'" + value
    return value


def _time(moment: Any) -> str | None:
    return to_ist(moment).strftime("%H:%M") if moment else None


def _delinquency_text(row: Row) -> str:
    return "; ".join(
        DELINQUENCY_LABELS[d["code"]] + (f" ({d['status'].lower()})" if d["status"] else "")
        for d in row.delinquencies
    )


@router.get("/attendance/report")
def attendance_report(
    request: Request,
    start: Annotated[str | None, Query(alias="from", max_length=10)] = None,
    end: Annotated[str | None, Query(alias="to", max_length=10)] = None,
    project_code: Annotated[str | None, Query(alias="projectCode", max_length=60)] = None,
    user: HumanPrincipal = Depends(can_read_attendance),
    clock: Clock = Depends(get_clock),
    conn: Connection = Depends(get_conn),
) -> Response:
    """An Excel report of check-in, check-out and delinquencies, one row per employee, per day,
    per project. One day by default; up to 31 days."""
    today = ist_date(clock())
    first = _day(start, today, "From")
    last = _day(end, first, "To")
    if last < first:
        raise ApiError(422, "HR_VALIDATION_FAILED", "The end date is before the start date.")
    if last > today:
        raise ApiError(422, "HR_VALIDATION_FAILED", "The report cannot include a future day.")
    if (last - first) >= timedelta(days=MAX_RANGE_DAYS):
        raise ApiError(
            422, "HR_VALIDATION_FAILED", f"Choose at most {MAX_RANGE_DAYS} days for one report."
        )
    rows = build_rows(conn, first, last, today=today, project_code=project_code)

    book = Workbook()
    sheet = book.active
    sheet.title = "Attendance"
    sheet.append(
        [
            "Date",
            "Employee ID",
            "Name",
            "Project",
            "Role",
            "Assigned outlets",
            "Status",
            "Check-in",
            "Check-out",
            "Hours",
            "Check-in outlet",
            "Check-in distance (m)",
            "Check-out outlet",
            "Check-out distance (m)",
            "Delinquencies",
        ]
    )
    for r in rows:
        sheet.append(
            [
                r.work_date.isoformat(),
                _text(r.employee_code),
                _text(r.employee_name),
                _text(r.project_name or "No project"),
                ", ".join(r.roles),
                _text(", ".join(r.outlets)),
                r.status.replace("_", " ").title(),
                _time(r.check_in_at),
                _time(r.check_out_at),
                r.hours,
                _text(r.check_in_outlet),
                r.check_in_distance_m,
                _text(r.check_out_outlet),
                r.check_out_distance_m,
                _delinquency_text(r),
            ]
        )
    late = book.create_sheet("Delinquencies")
    late.append(["Date", "Employee ID", "Name", "Project", "Delinquency", "Decision", "Reason"])
    for r in rows:
        for d in r.delinquencies:
            late.append(
                [
                    r.work_date.isoformat(),
                    _text(r.employee_code),
                    _text(r.employee_name),
                    _text(r.project_name or "No project"),
                    DELINQUENCY_LABELS[d["code"]],
                    d["status"].title() if d["status"] else "",
                    _text(d.get("reason")),
                ]
            )
    for ws in (sheet, late):
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
    out = io.BytesIO()
    book.save(out)

    record_audit(
        conn,
        actor_user_id=user.user_id,
        action="ATTENDANCE_REPORT_DOWNLOADED",
        entity_type="attendance_report",
        entity_id=f"{first.isoformat()}..{last.isoformat()}",
        changes={"projectCode": project_code, "rows": len(rows)},
        request=request,
    )
    name = (
        f"attendance-{first.isoformat()}.xlsx"
        if first == last
        else f"attendance-{first.isoformat()}-to-{last.isoformat()}.xlsx"
    )
    return Response(
        content=out.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="{name}"',
            "Cache-Control": "private, no-store",
        },
    )


@router.get("/work-assignments")
def work_assignments(
    project_code: Annotated[str | None, Query(alias="projectCode", max_length=60)] = None,
    employee_id: Annotated[str | None, Query(alias="employeeId", max_length=40)] = None,
    _: HumanPrincipal = Depends(can_read_employees),
    clock: Clock = Depends(get_clock),
    conn: Connection = Depends(get_conn),
) -> dict[str, Any]:
    """Which employee works on which projects, in which role and at which outlets, as copied
    from Audit Core. Read only: it is changed in Audit Core, then follows by the daily sync."""
    from hrmgmt.api.employees import _uuid

    now = clock()
    eid = _uuid(employee_id) if employee_id else None
    people = (
        conn.execute(
            text(
                """
                SELECT employee_id, employee_code, full_name, security_user_id, login_status
                FROM hr.employee
                WHERE employment_status = 'ACTIVE'
                  AND (CAST(:e AS uuid) IS NULL OR employee_id = CAST(:e AS uuid))
                ORDER BY lower(full_name), employee_code
                """
            ),
            {"e": eid},
        )
        .mappings()
        .all()
    )
    current: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for r in conn.execute(
        text(
            "SELECT security_user_id, project_code, project_name, role_code, dealer_name,"
            " outlet_name, latitude, valid_from FROM hr.work_assignment"
            " WHERE valid_from <= :now AND (valid_to IS NULL OR valid_to > :now)"
            " ORDER BY project_name, role_code, outlet_name"
        ),
        {"now": now},
    ).mappings():
        current[str(r["security_user_id"])].append(dict(r))
    projects: dict[str, str | None] = {}
    employees = []
    for p in people:
        items = current.get(str(p["security_user_id"]), []) if p["security_user_id"] else []
        for a in items:
            projects[a["project_code"]] = a["project_name"]
        if project_code:
            items = [a for a in items if a["project_code"] == project_code]
            if not items:
                continue
        employees.append(
            {
                "employeeId": str(p["employee_id"]),
                "employeeCode": p["employee_code"],
                "fullName": p["full_name"],
                "hasLogin": p["security_user_id"] is not None,
                "assignments": [
                    {
                        "projectCode": a["project_code"],
                        "projectName": a["project_name"],
                        "role": a["role_code"],
                        "dealerName": a["dealer_name"],
                        "outletName": a["outlet_name"],
                        "outletHasLocation": a["latitude"] is not None,
                        "since": a["valid_from"].isoformat(),
                    }
                    for a in items
                ],
            }
        )
    sync = wc.sync_status(conn)
    return {
        "syncedAt": sync["last_success_at"].isoformat() if sync.get("last_success_at") else None,
        "syncStatus": sync.get("last_status"),
        "projects": [
            {"projectCode": c, "projectName": n}
            for c, n in sorted(projects.items(), key=lambda kv: (kv[1] or "", kv[0] or ""))
        ],
        "employees": employees,
    }


def collapse_project_history(rows: list[dict[str, Any]], now: Any) -> list[dict[str, Any]]:
    """One line per project, role and outlet: the first start, and either still current or the last end.
    Audit Core closes and re-adds a person's rows whenever their mapping is edited, so the copy repeats."""
    groups: dict[tuple[Any, ...], dict[str, Any]] = {}
    for r in rows:
        current = r["valid_from"] <= now and (r["valid_to"] is None or r["valid_to"] > now)
        key = (r["tenant_id"], r["role_code"], r["outlet_id"])
        line = groups.get(key)
        if line is None:
            groups[key] = {
                "projectCode": r["project_code"],
                "projectName": r["project_name"],
                "role": r["role_code"],
                "dealerName": r["dealer_name"],
                "outletName": r["outlet_name"],
                "since": r["valid_from"],
                "until": r["valid_to"],
                "current": current,
            }
            continue
        line["since"] = min(line["since"], r["valid_from"])
        line["current"] = line["current"] or current
        if r["valid_to"] is not None and (line["until"] is None or r["valid_to"] > line["until"]):
            line["until"] = r["valid_to"]
    lines = list(groups.values())
    for line in lines:
        if line["current"]:
            line["until"] = None
    lines.sort(key=lambda line: (not line["current"], -line["since"].timestamp()))
    return lines


@router.get("/employees/{employee_id}/project-history")
def project_history(
    employee_id: str,
    _: HumanPrincipal = Depends(can_read_employees),
    clock: Clock = Depends(get_clock),
    conn: Connection = Depends(get_conn),
) -> dict[str, Any]:
    """Every project this person is, or has been, tagged to, with from and to dates, from the copy
    of Audit Core's assignments. Read only. Ends are as exact as the daily refresh."""
    from hrmgmt.api.employees import _uuid

    user = wc.employee_user_id(conn, _uuid(employee_id))
    if user is None:
        return {"linked": False, "syncedAt": None, "items": []}
    rows = (
        conn.execute(
            text(
                "SELECT tenant_id, project_code, project_name, role_code, dealer_name, outlet_id,"
                " outlet_name, valid_from, valid_to FROM hr.work_assignment"
                " WHERE security_user_id = CAST(:u AS uuid)"
            ),
            {"u": user},
        )
        .mappings()
        .all()
    )
    sync = wc.sync_status(conn)
    items = collapse_project_history([dict(r) for r in rows], clock())
    return {
        "linked": True,
        "syncedAt": sync["last_success_at"].isoformat() if sync.get("last_success_at") else None,
        "items": [
            {
                **{k: v for k, v in line.items() if k not in ("since", "until")},
                "since": line["since"].isoformat(),
                "until": line["until"].isoformat() if line["until"] else None,
            }
            for line in items
        ],
    }
