"""Bulk set designation and propose salary from an Excel sheet (Employee ID, Designation,
Salary (Monthly), Date of Joining). Preview first, nothing saved; then commit in batches.

A salary is only ever proposed here, exactly as one entered by hand: Finance still approves it.
A gross in the band that has no template (21,001 to 25,000) is not guessed; it is left for HR to
add on the employee's page, where the template is chosen on purpose."""

from __future__ import annotations

import hashlib
import io
import re
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Annotated, Any

import structlog
from fastapi import APIRouter, Depends, File, Form, Request, UploadFile
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import Connection, text

from hrmgmt import permissions as perm
from hrmgmt.api.employees import _uuid
from hrmgmt.api.payroll import BAND_HIGH, BAND_LOW, StructureIn, create_structure
from hrmgmt.audit import record_audit
from hrmgmt.db import get_conn
from hrmgmt.errors import ApiError
from hrmgmt.importer import MAX_IMPORT_BYTES, MAX_SHEET_ROWS
from hrmgmt.principal import current_user, get_authorizer, has_permission
from hrmgmt.security import HumanPrincipal

logger = structlog.get_logger(__name__)
router = APIRouter(prefix="/hr/v1", tags=["Designation and salary import"])

MAX_ROWS_PER_COMMIT = 10
_ALIASES = {
    "employee id": "code",
    "employee code": "code",
    "emp id": "code",
    "designation": "designation",
    "salary monthly": "salary",
    "monthly salary": "salary",
    "salary": "salary",
    "gross monthly": "salary",
    "date of joining": "joined",
    "doj": "joined",
}


def _header(raw: Any) -> str:
    return " ".join(re.sub(r"[^a-z0-9 ]+", " ", str(raw or "").lower()).split())


def _require_both(authorizer: Any, user: HumanPrincipal) -> None:
    for key in (perm.HR_EMPLOYEE_MANAGE, perm.HR_SALARY_PROPOSE):
        if not has_permission(authorizer, user, key):
            raise ApiError(403, "HR_FORBIDDEN", "You are not allowed to do this.")


async def _read(file: UploadFile) -> bytes:
    data = await file.read(MAX_IMPORT_BYTES + 1)
    if len(data) > MAX_IMPORT_BYTES:
        raise ApiError(422, "HR_IMPORT_FILE_INVALID", "The file is larger than 2 MB.")
    return data


def _parse(data: bytes) -> list[dict[str, Any]]:
    if not zipfile.is_zipfile(io.BytesIO(data)):
        raise ApiError(422, "HR_IMPORT_FILE_INVALID", "Upload an Excel .xlsx file.")
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many types for a damaged file
        raise ApiError(
            422, "HR_IMPORT_FILE_INVALID", "This file could not be read as an Excel workbook."
        ) from exc
    try:
        sheet = book[book.sheetnames[0]]
        columns: dict[int, str] | None = None
        rows: list[dict[str, Any]] = []
        for number, cells in enumerate(
            sheet.iter_rows(max_row=MAX_SHEET_ROWS, values_only=True), 1
        ):
            if columns is None:
                mapped = {
                    i: _ALIASES[_header(c)] for i, c in enumerate(cells) if _header(c) in _ALIASES
                }
                if {"code", "designation", "salary"} <= set(mapped.values()):
                    columns = mapped
                continue
            record = {columns[i]: cells[i] for i in columns if i < len(cells)}
            if all(v in (None, "") for v in record.values()):
                continue
            rows.append({"row": number, **record})
    finally:
        book.close()
    if columns is None:
        raise ApiError(
            422,
            "HR_IMPORT_FILE_INVALID",
            "No header row found. The sheet needs Employee ID, Designation and Salary (Monthly).",
        )
    return rows


def _joined(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def _money(value: Any) -> Decimal | None:
    try:
        amount = Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, ValueError):
        return None
    if amount <= 0 or amount > 10_000_000 or amount != amount.quantize(Decimal("0.01")):
        return None
    return amount


def _plan(conn: Connection, rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    designations = {
        str(r[1]).strip().lower(): str(r[0])
        for r in conn.execute(text("SELECT code, label FROM hr.designation WHERE active"))
    }
    codes = [str(r["code"] or "").strip().upper() for r in rows]
    people = {
        str(r["employee_code"]): r
        for r in conn.execute(
            text(
                "SELECT employee_id, employee_code, full_name, designation_code, date_of_joining,"
                " employment_status FROM hr.employee WHERE employee_code = ANY(:codes)"
            ),
            {"codes": codes},
        ).mappings()
    }
    plans: list[dict[str, Any]] = []
    seen: set[str] = set()
    for r, code in zip(rows, codes, strict=True):
        plan: dict[str, Any] = {
            "row": r["row"],
            "employeeCode": code,
            "errors": [],
            "notes": [],
            "setDesignation": None,
            "salary": None,
            "salaryAction": None,
        }
        person = people.get(code)
        if not code or person is None:
            plan["errors"].append("No employee has this Employee ID.")
        elif person["employment_status"] != "ACTIVE":
            plan["errors"].append("This employee is not active.")
        if code in seen:
            plan["errors"].append("Employee ID is repeated in this file.")
        seen.add(code)
        label = str(r["designation"] or "").strip()
        new_designation = designations.get(label.lower())
        if new_designation is None:
            plan["errors"].append(f"Designation {label!r} is not one of the listed designations.")
        amount = _money(r["salary"])
        if amount is None:
            plan["errors"].append("Salary is not a valid monthly amount.")
        if plan["errors"] or person is None:
            plans.append(plan)
            continue
        plan["employeeId"] = str(person["employee_id"])
        plan["fullName"] = person["full_name"]
        if person["designation_code"] != new_designation:
            plan["setDesignation"] = new_designation
        joined = _joined(r.get("joined"))
        plan["effectiveFrom"] = (joined or date.today()).isoformat()
        if joined and person["date_of_joining"] and joined != person["date_of_joining"]:
            plan["notes"].append("Date of joining differs from the record; the record is kept.")
        plan["salary"] = amount
        if BAND_LOW <= amount <= BAND_HIGH:
            plan["salaryAction"] = "NEEDS_TEMPLATE"
            plan["notes"].append(
                "This gross has no template; add the salary on the employee's page and choose one."
            )
        else:
            existing = conn.execute(
                text(
                    "SELECT status, gross_monthly FROM hr.salary_structure"
                    " WHERE employee_id = CAST(:e AS uuid)"
                    " AND (status = 'PROPOSED' OR (status = 'APPROVED' AND gross_monthly = :g))"
                ),
                {"e": plan["employeeId"], "g": amount},
            ).first()
            plan["salaryAction"] = "EXISTS" if existing else "PROPOSE"
            if existing:
                plan["notes"].append("A matching salary is already proposed or approved.")
        plans.append(plan)
    return plans


def _public(plan: dict[str, Any], names: dict[str, str]) -> dict[str, Any]:
    changes = plan["setDesignation"] is not None or plan["salaryAction"] == "PROPOSE"
    status = "ERROR" if plan["errors"] else ("READY" if changes else "NO_CHANGE")
    if not plan["errors"] and not changes and plan["salaryAction"] == "NEEDS_TEMPLATE":
        status = "NEEDS_TEMPLATE"
    return {
        "row": plan["row"],
        "status": status,
        "employeeCode": plan["employeeCode"],
        "fullName": plan.get("fullName"),
        "designation": names.get(plan["setDesignation"] or ""),
        "salaryAction": plan["salaryAction"],
        "salary": float(plan["salary"]) if plan["salary"] is not None else None,
        "effectiveFrom": plan.get("effectiveFrom"),
        "notes": plan["notes"],
        "errors": plan["errors"],
    }


def _names(conn: Connection) -> dict[str, str]:
    return {
        str(r[0]): str(r[1]) for r in conn.execute(text("SELECT code, label FROM hr.designation"))
    }


@router.post("/employees/designation-salary-import/preview")
async def preview(
    file: UploadFile = File(...),
    user: HumanPrincipal = Depends(current_user),
    authorizer: Any = Depends(get_authorizer),
    conn: Connection = Depends(get_conn),
) -> dict[str, Any]:
    """Say what would happen. Nothing is saved."""
    _require_both(authorizer, user)
    plans = _plan(conn, _parse(await _read(file)))
    names = _names(conn)
    out = [_public(p, names) for p in plans]
    count = {
        s: sum(1 for o in out if o["status"] == s)
        for s in ("READY", "NEEDS_TEMPLATE", "NO_CHANGE", "ERROR")
    }
    return {
        "summary": {
            "total": len(out),
            "ready": count["READY"],
            "needsTemplate": count["NEEDS_TEMPLATE"],
            "noChange": count["NO_CHANGE"],
            "errors": count["ERROR"],
        },
        "rows": out,
    }


@router.post("/employees/designation-salary-import/commit")
async def commit(
    request: Request,
    file: UploadFile = File(...),
    rows: Annotated[str, Form(max_length=200)] = "",
    user: HumanPrincipal = Depends(current_user),
    authorizer: Any = Depends(get_authorizer),
    conn: Connection = Depends(get_conn),
) -> dict[str, Any]:
    """Apply the listed rows (at most 10 per call). The sheet is read again here, so what is saved
    is what the sheet says now, checked again against what is on file."""
    _require_both(authorizer, user)
    try:
        wanted = sorted({int(x) for x in rows.split(",") if x.strip()})
    except ValueError as exc:
        raise ApiError(422, "HR_VALIDATION_FAILED", "Rows must be a list of row numbers.") from exc
    if not wanted or len(wanted) > MAX_ROWS_PER_COMMIT:
        raise ApiError(
            422,
            "HR_VALIDATION_FAILED",
            f"Choose between 1 and {MAX_ROWS_PER_COMMIT} rows to apply at a time.",
        )
    data = await _read(file)
    by_row = {p["row"]: p for p in _plan(conn, _parse(data))}
    results: list[dict[str, Any]] = []
    for number in wanted:
        plan = by_row.get(number)
        base = {"row": number, "employeeCode": plan["employeeCode"] if plan else None}
        if plan is None or plan["errors"]:
            results.append({**base, "status": "SKIPPED", "error": "Not ready to apply."})
            continue
        try:
            item = _apply(conn, plan, user, request)
            conn.commit()
        except (ApiError, ValidationError) as exc:
            conn.rollback()
            message = exc.detail if isinstance(exc, ApiError) else "Some values are not valid."
            results.append({**base, "status": "FAILED", "error": message})
            continue
        except Exception as exc:  # one bad row must not stop the others
            conn.rollback()
            logger.error(
                "hr_designation_salary_row_failed", row=number, error_type=type(exc).__name__
            )
            results.append({**base, "status": "FAILED", "error": "The row could not be saved."})
            continue
        results.append({**base, "status": "APPLIED", **item})
    record_audit(
        conn,
        actor_user_id=user.user_id,
        action="DESIGNATIONS_AND_SALARIES_IMPORTED",
        entity_type="designation_salary_import",
        entity_id=hashlib.sha256(data).hexdigest()[:16],
        changes={
            "requested": len(wanted),
            "applied": sum(1 for r in results if r["status"] == "APPLIED"),
            "failed": sum(1 for r in results if r["status"] == "FAILED"),
        },
        request=request,
    )
    return {"results": results}


def _apply(
    conn: Connection, plan: dict[str, Any], user: HumanPrincipal, request: Request
) -> dict[str, Any]:
    eid = _uuid(plan["employeeId"])
    item: dict[str, Any] = {"designation": "UNCHANGED", "salary": plan["salaryAction"]}
    if plan["setDesignation"] is not None:
        before = conn.execute(
            text(
                "SELECT designation_code FROM hr.employee WHERE employee_id = CAST(:e AS uuid)"
                " FOR UPDATE"
            ),
            {"e": eid},
        ).scalar_one()
        conn.execute(
            text(
                "UPDATE hr.employee SET designation_code = :d, updated_at = now(), updated_by = :u"
                " WHERE employee_id = CAST(:e AS uuid)"
            ),
            {"d": plan["setDesignation"], "u": user.user_id, "e": eid},
        )
        record_audit(
            conn,
            actor_user_id=user.user_id,
            action="EMPLOYEE_UPDATED",
            entity_type="employee",
            entity_id=eid,
            changes={"designation_code": {"from": before, "to": plan["setDesignation"]}},
            request=request,
        )
        item["designation"] = "UPDATED"
    if plan["salaryAction"] == "PROPOSE":
        create_structure(
            conn,
            StructureIn(
                employee_id=eid,
                gross_monthly=plan["salary"],
                effective_from=date.fromisoformat(plan["effectiveFrom"]),
                note="Imported from the designation and salary sheet",
            ),
            user,
            request,
        )
        item["salary"] = "PROPOSED"
    return item
