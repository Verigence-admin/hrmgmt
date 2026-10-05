from __future__ import annotations

import uuid
from pathlib import Path

import numpy as np
from PIL import Image, ImageDraw, ImageEnhance
from sqlalchemy import text

from hrmgmt import permissions as perm
from hrmgmt.facecheck import face_present
from tests.support import NEAR, OUTLET, World, ist
from tests.test_attendance import send, token

# A public-domain NASA portrait (scikit-image's "astronaut"), only to prove the detector works at all.
_FACE = Path(__file__).parent / "fixtures" / "face_public_domain.jpg"


def test_a_plain_picture_has_no_face():
    assert face_present(Image.new("RGB", (640, 480), (120, 160, 200))) is False


def test_a_real_face_is_found_even_when_dim_tilted_or_small():
    face = Image.open(_FACE).convert("RGB")
    assert face_present(face) is True
    assert face_present(ImageEnhance.Brightness(face).enhance(0.6)) is True
    assert face_present(face.rotate(8, fillcolor=(128, 128, 128))) is True
    canvas = Image.new("RGB", (640, 480), (150, 150, 150))  # the same face, farther from the camera
    canvas.paste(face.resize((300, 300)), (170, 90))
    assert face_present(canvas) is True


def test_a_tilted_or_partly_cut_off_face_is_still_a_face():
    # What the older detector alone missed: a face tilted well over, and one half out of the frame.
    face = Image.open(_FACE).convert("RGB")
    width, height = face.size
    assert face_present(face.rotate(-20, fillcolor=(128, 128, 128))) is True
    assert face_present(face.crop((int(width * 0.38), 0, width, height))) is True
    assert face_present(face.crop((int(width * 0.45), 0, width, height))) is True


def test_if_the_new_detector_is_missing_the_older_one_still_runs(monkeypatch):
    from hrmgmt import facecheck

    monkeypatch.setattr(facecheck, "_yunet_ready", lambda: False)
    assert face_present(Image.open(_FACE).convert("RGB")) is True
    assert face_present(Image.new("RGB", (640, 480), (120, 160, 200))) is False


def test_if_no_detector_can_run_nothing_is_flagged(monkeypatch):
    from hrmgmt import facecheck

    monkeypatch.setattr(facecheck, "_yunet_ready", lambda: False)
    monkeypatch.setattr(facecheck, "_detector", lambda: None)
    assert face_present(Image.new("RGB", (640, 480), (120, 160, 200))) is None


def test_a_ceiling_a_wall_and_noise_have_no_face():
    ceiling = Image.new("RGB", (640, 480), (225, 225, 220))
    draw = ImageDraw.Draw(ceiling)
    draw.ellipse(
        (250, 120, 370, 240), fill=(250, 250, 245), outline=(190, 190, 185), width=6
    )  # a ceiling light
    draw.ellipse((470, 40, 510, 80), fill=(150, 150, 150))  # a smoke detector
    noise = Image.fromarray(
        np.random.default_rng(7).integers(0, 255, (480, 640, 3), dtype=np.uint8)
    )
    stripes = Image.fromarray(np.tile(np.arange(640) % 40 * 6, (480, 1)).astype(np.uint8)).convert(
        "RGB"
    )
    for picture in (ceiling, noise, stripes):
        assert face_present(picture) is False


def test_a_photo_without_a_face_is_accepted_but_flagged_for_hr(migrated_engine, monkeypatch):
    world = World(migrated_engine, now=ist(2026, 10, 5, 10, 0))
    world.clean_assignments()
    with migrated_engine.begin() as conn:
        for table in ("attendance_exception", "attendance_day", "capture_token"):
            conn.execute(text(f"DELETE FROM hr.{table}"))
    hr = world.grant(str(uuid.uuid4()), perm.HR_EMPLOYEE_MANAGE)
    keeper = world.grant(str(uuid.uuid4()), perm.HR_ATTENDANCE_READ_ALL)
    emp, user = world.employee(hr)
    world.assign(user, "PC", outlet=OUTLET)
    monkeypatch.setattr("hrmgmt.api.attendance.face_present", lambda image: False)
    sent = send(world, user, "in", token(world, user), where=NEAR)
    assert sent.status_code == 200, sent.text
    assert sent.json()["flags"] == ["NO_FACE"] and sent.json()["needsApproval"] == []
    daily = world.client.get(
        "/hr/v1/attendance/daily", params={"date": "2026-10-05"}, headers=world.headers(keeper)
    ).json()
    mine = [r for r in daily["rows"] if r["employeeCode"] == emp["employeeCode"]][0]
    assert [d["code"] for d in mine["delinquencies"]] == ["NO_FACE"]
    assert mine["delinquencies"][0]["label"] == "No face found in the check-in photo"
    assert mine["status"] == "CHECKED_IN"  # nothing waits for approval


def test_the_flag_says_which_photo_it_is_for(migrated_engine, monkeypatch):
    world = World(migrated_engine, now=ist(2026, 10, 5, 10, 0))
    world.clean_assignments()
    with migrated_engine.begin() as conn:
        for table in ("attendance_exception", "attendance_day", "capture_token"):
            conn.execute(text(f"DELETE FROM hr.{table}"))
    hr = world.grant(str(uuid.uuid4()), perm.HR_EMPLOYEE_MANAGE)
    keeper = world.grant(str(uuid.uuid4()), perm.HR_ATTENDANCE_READ_ALL)
    only_out, out_user = world.employee(hr)
    only_in, in_user = world.employee(hr)
    both, both_user = world.employee(hr)
    for user in (out_user, in_user, both_user):
        world.assign(user, "PC", outlet=OUTLET)
    faces = {"in": {in_user: False, both_user: False}, "out": {out_user: False, both_user: False}}
    current: dict[str, object] = {}
    monkeypatch.setattr(
        "hrmgmt.api.attendance.face_present",
        lambda image: faces[current["event"]].get(current["user"], True),
    )

    def punch(user: str, event: str, purpose: str):
        current.update(user=user, event=event)
        return send(world, user, event, token(world, user, purpose), where=NEAR)

    for user in (out_user, in_user, both_user):
        assert punch(user, "in", "CHECK_IN").status_code == 200
    world.clock.set(ist(2026, 10, 5, 18, 30))
    for user in (out_user, in_user, both_user):
        assert punch(user, "out", "CHECK_OUT").status_code == 200

    rows = {
        r["employeeCode"]: r
        for r in world.client.get(
            "/hr/v1/attendance/daily", params={"date": "2026-10-05"}, headers=world.headers(keeper)
        ).json()["rows"]
    }
    seen = {
        name: [(d["label"], d["side"]) for d in rows[e["employeeCode"]]["delinquencies"]]
        for name, e in (("out", only_out), ("in", only_in), ("both", both))
    }
    assert seen == {
        "out": [("No face found in the check-out photo", "CHECK_OUT")],
        "in": [("No face found in the check-in photo", "CHECK_IN")],
        "both": [
            ("No face found in the check-in photo", "CHECK_IN"),
            ("No face found in the check-out photo", "CHECK_OUT"),
        ],
    }
    # the same wording reaches the Excel report
    world.clock.set(ist(2026, 10, 6, 11, 0))
    report = world.client.get(
        "/hr/v1/attendance/report",
        params={"from": "2026-10-05", "to": "2026-10-05"},
        headers=world.headers(keeper),
    )
    import io

    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(report.content))
    texts = {
        row[1]: row[14]
        for row in book["Attendance"].iter_rows(min_row=2, values_only=True)
        if row[1] in {only_out["employeeCode"], only_in["employeeCode"], both["employeeCode"]}
    }
    assert texts[only_out["employeeCode"]] == "No face found in the check-out photo"
    assert texts[only_in["employeeCode"]] == "No face found in the check-in photo"
    assert (
        "check-in photo" in texts[both["employeeCode"]]
        and "check-out photo" in texts[both["employeeCode"]]
    )
    listed = [
        (r[4], r[0 + 1])
        for r in book["Delinquencies"].iter_rows(min_row=2, values_only=True)
        if r[1] == only_out["employeeCode"]
    ]
    assert listed == [("No face found in the check-out photo", only_out["employeeCode"])]


def test_when_the_check_cannot_run_nothing_is_flagged(migrated_engine, monkeypatch):
    world = World(migrated_engine, now=ist(2026, 10, 5, 10, 0))
    world.clean_assignments()
    with migrated_engine.begin() as conn:
        for table in ("attendance_exception", "attendance_day", "capture_token"):
            conn.execute(text(f"DELETE FROM hr.{table}"))
    hr = world.grant(str(uuid.uuid4()), perm.HR_EMPLOYEE_MANAGE)
    _, user = world.employee(hr)
    world.assign(user, "PC", outlet=OUTLET)
    monkeypatch.setattr("hrmgmt.api.attendance.face_present", lambda image: None)
    sent = send(world, user, "in", token(world, user), where=NEAR)
    assert sent.status_code == 200 and sent.json()["flags"] == []
