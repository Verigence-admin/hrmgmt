from __future__ import annotations

import io
import uuid
from datetime import timedelta

import pytest
from openpyxl import load_workbook
from sqlalchemy import text

from hrmgmt import permissions as perm
from tests.support import FAR, NEAR, OUTLET, World, ist
from tests.test_attendance import send, token

SECOND = ("Bhubaneswar Motors", 20.2961, 85.8245)


@pytest.fixture()
def world(migrated_engine):
    w = World(migrated_engine)
    w.clean_assignments()
    with migrated_engine.begin() as conn:
        for table in (
            "attendance_exception",
            "attendance_day",
            "capture_token",
        ):
            conn.execute(text(f"DELETE FROM hr.{table}"))
        conn.execute(text("DELETE FROM hr.leave_request"))
        conn.execute(text("DELETE FROM hr.holiday WHERE status = 'DECLARED'"))
        conn.execute(text("DELETE FROM hr.setting"))
    return w


def _hr(world: World) -> tuple[str, str]:
    admin = world.grant(str(uuid.uuid4()), perm.HR_EMPLOYEE_MANAGE)
    keeper = world.grant(str(uuid.uuid4()), perm.HR_ATTENDANCE_READ_ALL, perm.HR_EMPLOYEE_READ)
    return admin, keeper


def _person(world: World, admin: str, **assign):
    emp, user = world.employee(admin, date_of_joining="2026-01-05")
    return emp, user


def _rows(world: World, keeper: str, day: str, **params) -> dict[str, list[dict]]:
    r = world.client.get(
        "/hr/v1/attendance/daily", params={"date": day, **params}, headers=world.headers(keeper)
    )
    assert r.status_code == 200, r.text
    out: dict[str, list[dict]] = {}
    for row in r.json()["rows"]:
        out.setdefault(row["employeeCode"], []).append(row)
    return out


def test_daily_view_shows_who_is_in_out_absent_or_on_leave_and_what_is_wrong(
    world, migrated_engine
):
    admin, keeper = _hr(world)
    done, done_user = _person(world, admin)
    open_day, open_user = _person(world, admin)
    absent, absent_user = _person(world, admin)
    away, away_user = _person(world, admin)
    for u in (done_user, open_user, absent_user, away_user):
        world.assign(u, "PC", outlet=OUTLET)
    with migrated_engine.begin() as conn:
        conn.execute(
            text(
                "INSERT INTO hr.leave_request (employee_id, leave_type, from_date, to_date, days,"
                " status, approver_rule) VALUES (CAST(:e AS uuid), 'SICK', '2026-10-05',"
                " '2026-10-05', 1, 'APPROVED', 'HR')"
            ),
            {"e": away["employeeId"]},
        )
    send(world, done_user, "in", token(world, done_user), where=NEAR)
    world.clock.set(ist(2026, 10, 5, 18, 30))
    send(world, done_user, "out", token(world, done_user, "CHECK_OUT"), where=NEAR)
    world.clock.set(ist(2026, 10, 5, 10, 25))
    send(world, open_user, "in", token(world, open_user), where=NEAR)

    # the same day: still in progress, nobody is delinquent yet
    today = _rows(world, keeper, "2026-10-05")
    assert today[done["employeeCode"]][0]["status"] == "COMPLETE"
    assert today[done["employeeCode"]][0]["hoursWorked"] > 7
    assert today[open_day["employeeCode"]][0]["status"] == "CHECKED_IN"
    assert today[absent["employeeCode"]][0]["status"] == "NOT_CHECKED_IN"
    assert today[away["employeeCode"]][0]["status"] == "ON_LEAVE"
    assert today[absent["employeeCode"]][0]["delinquencies"] == []

    # the next day, that day is past: missing check-out and absence are now delinquencies
    world.clock.set(ist(2026, 10, 6, 11, 0))
    past = _rows(world, keeper, "2026-10-05")
    codes = {k: [d["code"] for d in v[0]["delinquencies"]] for k, v in past.items()}
    assert codes[open_day["employeeCode"]] == ["MISSING_CHECK_OUT"]
    assert codes[absent["employeeCode"]] == ["ABSENT"]
    assert codes[done["employeeCode"]] == [] and codes[away["employeeCode"]] == []
    body = world.client.get(
        "/hr/v1/attendance/daily", params={"date": "2026-10-05"}, headers=world.headers(keeper)
    ).json()
    assert body["dayKind"] == "WORKING" and body["summary"]["absent"] >= 1


def test_daily_view_never_hides_someone_who_checked_in_because_of_a_later_joining_date(
    world, migrated_engine
):
    admin, keeper = _hr(world)
    worker, worker_user = _person(world, admin)
    not_yet, not_yet_user = _person(world, admin)
    for u in (worker_user, not_yet_user):
        world.assign(u, "PC", outlet=OUTLET)
    send(world, worker_user, "in", token(world, worker_user), where=NEAR)
    with migrated_engine.begin() as conn:
        # a joining date typed in wrongly, later than the day they were at work
        conn.execute(
            text(
                "UPDATE hr.employee SET date_of_joining = '2026-10-14'"
                " WHERE employee_id IN (CAST(:a AS uuid), CAST(:b AS uuid))"
            ),
            {"a": worker["employeeId"], "b": not_yet["employeeId"]},
        )
    day = world.clock().date().isoformat()
    rows = _rows(world, keeper, day)
    # the person who checked in is listed, with the check-in
    assert rows[worker["employeeCode"]][0]["status"] == "CHECKED_IN"
    assert rows[worker["employeeCode"]][0]["checkInAt"] is not None
    # the person with no check-in and a later joining date is still not expected yet
    assert not_yet["employeeCode"] not in rows


def test_daily_view_carries_the_photo_flags_and_the_reason_for_being_away(world):
    admin, keeper = _hr(world)
    emp, user = _person(world, admin)
    world.assign(user, "PC", outlet=OUTLET)
    sent = send(
        world, user, "in", token(world, user), where=FAR, reason="Visiting the other showroom"
    )
    assert sent.status_code == 200, sent.text
    row = _rows(world, keeper, "2026-10-05")[emp["employeeCode"]][0]
    assert row["attendanceId"] == sent.json()["attendanceId"]
    assert row["hasCheckInPhoto"] is True and row["hasCheckOutPhoto"] is False
    away = [d for d in row["delinquencies"] if d["code"] == "OUT_OF_FENCE"]
    assert away and away[0]["reason"] == "Visiting the other showroom"
    assert away[0]["label"] == "Not in tagged location at check-in"
    # a day with no check-in has no photo and no attendance record
    other, other_user = _person(world, admin)
    world.assign(other_user, "PC", outlet=OUTLET)
    quiet = _rows(world, keeper, "2026-10-05")[other["employeeCode"]][0]
    assert quiet["attendanceId"] is None and quiet["hasCheckInPhoto"] is False


def test_hr_can_decide_a_day_from_the_daily_view_but_the_lists_stay_with_the_approvers(
    world, migrated_engine
):
    admin, keeper = _hr(world)
    emp, user = _person(world, admin)
    world.assign(user, "PC", outlet=OUTLET)
    _, tl_user = _person(world, admin)
    world.assign(tl_user, "TL")
    send(world, user, "in", token(world, user), where=FAR, reason="Other showroom")
    with migrated_engine.connect() as conn:
        xid = str(
            conn.execute(text("SELECT exception_id FROM hr.attendance_exception")).scalar_one()
        )
    # HR's own approvals list still holds only what HR is the approver for
    listed = world.client.get("/hr/v1/approvals/attendance", headers=world.headers(keeper))
    assert listed.json()["items"] == []
    # ...but HR, acting on the person's day, can decide it
    done = world.client.post(
        f"/hr/v1/approvals/attendance/{xid}/decision",
        json={"decision": "APPROVE"},
        headers=world.headers(keeper),
    )
    assert done.status_code == 200 and done.json()["status"] == "APPROVED"
    # someone without the HR attendance permission still cannot
    stranger = world.grant(str(uuid.uuid4()), perm.HR_EMPLOYEE_READ)
    refused = world.client.post(
        f"/hr/v1/approvals/attendance/{xid}/decision",
        json={"decision": "APPROVE"},
        headers=world.headers(stranger),
    )
    assert refused.status_code == 404


def test_hr_cannot_decide_their_own_exception(world, migrated_engine):
    admin, keeper = _hr(world)
    emp, user = _person(world, admin)
    world.assign(user, "PC", outlet=OUTLET)
    world.grant(user, perm.HR_ATTENDANCE_READ_ALL)
    send(world, user, "in", token(world, user), where=FAR, reason="Other showroom")
    with migrated_engine.connect() as conn:
        xid = str(
            conn.execute(text("SELECT exception_id FROM hr.attendance_exception")).scalar_one()
        )
    own = world.client.post(
        f"/hr/v1/approvals/attendance/{xid}/decision",
        json={"decision": "APPROVE"},
        headers=world.headers(user),
    )
    assert own.status_code == 404


def test_people_are_listed_alphabetically_whatever_their_codes(world):
    admin, keeper = _hr(world)
    for name in ("Chandini Nayak", "akash Das", "Bina Rout"):
        _, user = world.employee(admin, full_name=name)
        world.assign(user, "PC", outlet=OUTLET, project=("PA", "Project A"))
    rows = world.client.get(
        "/hr/v1/attendance/daily",
        params={"date": "2026-10-05", "projectCode": "PA"},
        headers=world.headers(keeper),
    ).json()["rows"]
    assert [r["employeeName"] for r in rows] == ["akash Das", "Bina Rout", "Chandini Nayak"]
    listed = world.client.get(
        "/hr/v1/employees", params={"limit": 100}, headers=world.headers(admin)
    )
    names = [e["fullName"] for e in listed.json()["items"]]
    assert names == sorted(names, key=str.lower)


def test_team_month_shows_working_days_holidays_leave_and_absence(world, migrated_engine):
    admin, keeper = _hr(world)
    world.grant(admin, perm.HR_SETTINGS_MANAGE)
    emp, user = _person(world, admin)
    world.assign(user, "PC", outlet=OUTLET)
    declared = world.client.put(
        "/hr/v1/holidays/2026-10-20",
        json={"name": "Vijaya Dasami", "status": "DECLARED"},
        headers=world.headers(admin),
    )
    assert declared.status_code == 200
    for day in (5, 6, 11):  # two working days and one Sunday
        world.clock.set(ist(2026, 10, day, 10, 20))
        assert send(world, user, "in", token(world, user)).status_code == 200
    with migrated_engine.begin() as conn:
        conn.execute(
            text(
                "INSERT INTO hr.leave_request (employee_id, leave_type, from_date, to_date, days,"
                " status, approver_rule) VALUES (CAST(:e AS uuid), 'SICK', '2026-10-07',"
                " '2026-10-07', 1, 'APPROVED', 'HR')"
            ),
            {"e": emp["employeeId"]},
        )
    world.clock.set(ist(2026, 10, 21, 10, 20))
    body = world.client.get(
        "/hr/v1/attendance/team", params={"month": "2026-10"}, headers=world.headers(keeper)
    ).json()
    assert body["summary"]["sundays"] == 4
    assert body["summary"]["holidays"] == [{"date": "2026-10-20", "name": "Vijaya Dasami"}]
    assert body["summary"]["workingDays"] == 26 and body["summary"]["workingDaysSoFar"] == 17
    mine = [e for e in body["employees"] if e["employeeCode"] == emp["employeeCode"]][0]
    assert mine["daysPresent"] == 2 and mine["offDayWorked"] == 1 and mine["daysOnLeave"] == 1
    # 16 working days have passed before today (21st); 2 present, 1 on leave
    assert mine["daysAbsent"] == 13
    assert mine["pendingExceptions"] == 1  # the Sunday


def test_a_person_on_two_projects_appears_once_for_each_and_can_be_filtered(world):
    admin, keeper = _hr(world)
    emp, user = _person(world, admin)
    world.assign(user, "PC", outlet=OUTLET, project=("P1", "Project One"))
    world.assign(user, "PC", outlet=SECOND, tenant="tenant-b", project=("P2", "Project Two"))
    rows = _rows(world, keeper, "2026-10-05")[emp["employeeCode"]]
    assert sorted(r["projectName"] for r in rows) == ["Project One", "Project Two"]
    only = _rows(world, keeper, "2026-10-05", projectCode="P2")
    assert [r["projectName"] for r in only[emp["employeeCode"]]] == ["Project Two"]


def test_sunday_and_declared_holiday_are_not_working_days(world):
    admin, keeper = _hr(world)
    emp, user = _person(world, admin)
    world.assign(user, "PC", outlet=OUTLET)
    world.clock.set(ist(2026, 10, 12, 11, 0))
    assert emp["employeeCode"] not in _rows(world, keeper, "2026-10-11")  # a Sunday
    r = world.client.get(
        "/hr/v1/attendance/daily", params={"date": "2026-10-11"}, headers=world.headers(keeper)
    )
    assert r.json()["dayKind"] == "SUNDAY"


def test_the_report_is_an_excel_file_with_every_row_and_the_delinquencies(world, migrated_engine):
    admin, keeper = _hr(world)
    emp, user = _person(world, admin)
    world.assign(user, "PC", outlet=OUTLET)
    send(world, user, "in", token(world, user), where=NEAR)
    world.clock.set(ist(2026, 10, 6, 11, 0))
    r = world.client.get(
        "/hr/v1/attendance/report",
        params={"from": "2026-10-05", "to": "2026-10-05"},
        headers=world.headers(keeper),
    )
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]
    assert "attendance-2026-10-05.xlsx" in r.headers["content-disposition"]
    book = load_workbook(io.BytesIO(r.content))
    assert book.sheetnames == ["Attendance", "Delinquencies"]
    mine = [
        row
        for row in book["Attendance"].iter_rows(values_only=True)
        if row[1] == emp["employeeCode"]
    ]
    assert len(mine) == 1 and mine[0][3] == "Project One" and mine[0][7] is not None
    assert "never checked out" in mine[0][14]
    bad = [
        row
        for row in book["Delinquencies"].iter_rows(values_only=True)
        if row[1] == emp["employeeCode"]
    ]
    assert [b[4] for b in bad] == ["Checked in, never checked out"]
    with migrated_engine.connect() as conn:
        assert (
            conn.execute(
                text(
                    "SELECT count(*) FROM hr.audit_log WHERE action = 'ATTENDANCE_REPORT_DOWNLOADED'"
                )
            ).scalar_one()
            >= 1
        )


@pytest.mark.parametrize(
    "params",
    [
        {"from": "2026-10-05", "to": "2026-12-30"},
        {"from": "2026-10-05", "to": "2026-10-04"},
        {"from": "2099-01-01"},
        {"from": "not-a-date"},
    ],
)
def test_the_report_refuses_bad_ranges(world, params):
    _, keeper = _hr(world)
    r = world.client.get("/hr/v1/attendance/report", params=params, headers=world.headers(keeper))
    assert r.status_code == 422


def test_only_hr_can_see_the_daily_view_and_the_report(world):
    admin, _ = _hr(world)
    nobody = world.grant(str(uuid.uuid4()))
    for path in ("/hr/v1/attendance/daily", "/hr/v1/attendance/report", "/hr/v1/work-assignments"):
        assert world.client.get(path, headers=world.headers(nobody)).status_code == 403
    assert (
        world.client.get("/hr/v1/work-assignments", headers=world.headers(admin)).status_code == 403
    )


def test_work_assignments_list_each_employees_projects_roles_and_outlets(world):
    admin, keeper = _hr(world)
    emp, user = _person(world, admin)
    lone, _ = _person(world, admin)
    world.assign(user, "PC", outlet=OUTLET, project=("P1", "Project One"))
    world.assign(
        user,
        "PC",
        outlet=("Unmapped", None, None),
        tenant="tenant-b",
        project=("P2", "Project Two"),
    )
    body = world.client.get("/hr/v1/work-assignments", headers=world.headers(keeper)).json()
    people = {e["employeeCode"]: e for e in body["employees"]}
    mine = people[emp["employeeCode"]]["assignments"]
    assert sorted((a["projectCode"], a["role"], a["outletHasLocation"]) for a in mine) == [
        ("P1", "PC", True),
        ("P2", "PC", False),
    ]
    assert people[lone["employeeCode"]]["assignments"] == []
    assert {p["projectCode"] for p in body["projects"]} >= {"P1", "P2"}
    only = world.client.get(
        "/hr/v1/work-assignments", params={"projectCode": "P2"}, headers=world.headers(keeper)
    ).json()
    assert [e["employeeCode"] for e in only["employees"]] == [emp["employeeCode"]]


def test_the_employee_list_names_each_persons_current_projects(world):
    admin, keeper = _hr(world)
    emp, user = _person(world, admin)
    lone, _ = _person(world, admin)
    world.assign(user, "PC", outlet=OUTLET, project=("P1", "Project One"))
    world.assign(
        user,
        "PC",
        outlet=("Unmapped", None, None),
        tenant="tenant-b",
        project=("P2", "Project Two"),
    )
    people = {}
    for code in (emp["employeeCode"], lone["employeeCode"]):  # by code: the full list is paged
        body = world.client.get(
            "/hr/v1/employees", params={"q": code}, headers=world.headers(keeper)
        ).json()
        people.update({e["employeeCode"]: e for e in body["items"]})
    assert people[emp["employeeCode"]]["projects"] == ["Project One", "Project Two"]
    assert people[lone["employeeCode"]]["projects"] == []


def test_project_history_folds_repeats_and_keeps_ended_projects(world, migrated_engine):
    admin, keeper = _hr(world)
    emp, user = _person(world, admin)
    now = world.clock()
    # Project One: closed by an edit and live again; Project Two: worked on, then ended.
    world.assign(
        user,
        "PC",
        outlet=OUTLET,
        project=("P1", "Project One"),
        valid_from=now - timedelta(days=30),
    )
    world.assign(
        user,
        "PC",
        outlet=OUTLET,
        project=("P2", "Project Two"),
        tenant="tenant-b",
        valid_from=now - timedelta(days=90),
    )
    with migrated_engine.begin() as conn:
        conn.execute(
            text("UPDATE hr.work_assignment SET valid_to = :t WHERE project_code = 'P2'"),
            {"t": now - timedelta(days=40)},
        )
    body = world.client.get(
        f"/hr/v1/employees/{emp['employeeId']}/project-history", headers=world.headers(keeper)
    ).json()
    assert body["linked"] is True
    assert [(i["projectName"], i["current"], i["until"] is None) for i in body["items"]] == [
        ("Project One", True, True),
        ("Project Two", False, False),
    ]
    assert (
        body["items"][0]["outletName"] == "Cuttack Motors"
        and body["items"][1]["since"] < body["items"][0]["since"]
    )


def test_project_history_needs_employee_read_and_says_when_there_is_no_login(world):
    admin, keeper = _hr(world)
    emp, _ = _person(world, admin)
    nobody = world.grant(str(uuid.uuid4()))
    assert (
        world.client.get(
            f"/hr/v1/employees/{emp['employeeId']}/project-history", headers=world.headers(nobody)
        ).status_code
        == 403
    )
    with world.engine.begin() as conn:
        conn.execute(
            text(
                "UPDATE hr.employee SET security_user_id = NULL WHERE employee_id = CAST(:e AS uuid)"
            ),
            {"e": emp["employeeId"]},
        )
    body = world.client.get(
        f"/hr/v1/employees/{emp['employeeId']}/project-history", headers=world.headers(keeper)
    ).json()
    assert body == {"linked": False, "syncedAt": None, "items": []}


def test_folding_repeated_rows_keeps_the_first_start_and_the_last_end():
    from datetime import UTC, datetime

    from hrmgmt.api.attendance_reports import collapse_project_history

    now = datetime(2026, 10, 5, tzinfo=UTC)

    def row(start, end):
        return {
            "tenant_id": "t",
            "project_code": "P",
            "project_name": "Project",
            "role_code": "PC",
            "dealer_name": "D",
            "outlet_id": "o",
            "outlet_name": "O",
            "valid_from": now + timedelta(days=start),
            "valid_to": None if end is None else now + timedelta(days=end),
        }

    [ended] = collapse_project_history([row(-30, -10), row(-10, -3)], now)
    assert (
        ended["current"] is False
        and ended["since"] == now - timedelta(days=30)
        and ended["until"] == now - timedelta(days=3)
    )
    [live] = collapse_project_history([row(-30, -10), row(-10, None)], now)
    assert live["current"] is True and live["until"] is None


def test_each_row_says_whether_a_check_in_or_out_was_outside_the_tagged_location(
    world, migrated_engine
):
    admin, keeper = _hr(world)
    inside, inside_user = _person(world, admin)
    outside, outside_user = _person(world, admin)
    no_outlet_position, no_position_user = _person(world, admin)
    lead, lead_user = _person(world, admin)
    absent, absent_user = _person(world, admin)
    world.assign(inside_user, "PC", outlet=OUTLET)
    world.assign(outside_user, "PC", outlet=OUTLET)
    world.assign(no_position_user, "PC", outlet=("No-GPS Motors", None, None))
    world.assign(lead_user, "TL", outlet=OUTLET)
    world.assign(absent_user, "PC", outlet=OUTLET)

    send(world, inside_user, "in", token(world, inside_user), where=NEAR)
    send(
        world,
        outside_user,
        "in",
        token(world, outside_user),
        where=FAR,
        reason="At the other showroom",
    )
    send(world, no_position_user, "in", token(world, no_position_user), where=NEAR)
    send(world, lead_user, "in", token(world, lead_user), where=FAR)
    world.clock.set(ist(2026, 10, 5, 18, 30))
    send(
        world,
        inside_user,
        "out",
        token(world, inside_user, "CHECK_OUT"),
        where=FAR,
        reason="Delivery",
    )

    rows = _rows(world, keeper, "2026-10-05")
    seen = {
        name: (
            rows[e["employeeCode"]][0]["checkInOutOfFence"],
            rows[e["employeeCode"]][0]["checkOutOutOfFence"],
        )
        for name, e in (
            ("inside", inside),
            ("outside", outside),
            ("no position", no_outlet_position),
            ("lead", lead),
            ("absent", absent),
        )
    }
    assert seen == {
        "inside": (False, True),  # in at the outlet, out somewhere else
        "outside": (True, None),  # in away from it; no check-out yet
        "no position": (None, None),  # the outlet has no location on file: nothing to say
        "lead": (None, None),  # the fence does not apply to a Team Lead
        "absent": (None, None),  # no punch
    }

    # the same values reach the Excel report, in two new columns at the far right
    world.clock.set(ist(2026, 10, 6, 11, 0))
    r = world.client.get(
        "/hr/v1/attendance/report",
        params={"from": "2026-10-05", "to": "2026-10-05"},
        headers=world.headers(keeper),
    )
    sheet = load_workbook(io.BytesIO(r.content))["Attendance"]
    header = [c.value for c in sheet[1]]
    assert header[-2:] == ["Check-in out of fence", "Check-out out of fence"]
    assert header[:15] == [
        "Date",
        "Employee ID",
        "Name",
        "Project",
        "Role",
        "Assigned outlets",
        "Status",
        "Check-in",
        "Check-out",
        "Hours",
        "Check-in outlet",
        "Check-in distance (m)",
        "Check-out outlet",
        "Check-out distance (m)",
        "Delinquencies",
    ]  # nothing that was already there has moved
    cells = {row[1]: row[-2:] for row in sheet.iter_rows(min_row=2, values_only=True)}
    assert cells[inside["employeeCode"]] == ("No", "Yes")
    assert cells[outside["employeeCode"]] == ("Yes", None)
    assert cells[no_outlet_position["employeeCode"]] == (None, None)
    assert cells[lead["employeeCode"]] == (None, None)
